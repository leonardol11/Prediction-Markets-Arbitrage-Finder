"""Write sheet models (see sheets.py) to a styled .xlsx workbook."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="111827")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
TITLE_FONT = Font(bold=True, size=11, name="Calibri", color="111827")
SUB_FONT = Font(size=9, name="Calibri", color="6B7280")
BODY_FONT = Font(size=10, name="Calibri")
LINK_FONT = Font(size=10, name="Calibri", color="1D4ED8", underline="single")
POS_FONT = Font(size=10, name="Calibri", color="047857", bold=True)
BOTTOM = Border(bottom=Side(style="thin", color="E5E7EB"))
BAND = PatternFill("solid", fgColor="F9FAFB")

ALIGN = {"left": "left", "center": "center", "right": "right"}


def _unpack(c):
    """cell → (value, display, url, number_format)."""
    if isinstance(c, dict):
        return c.get("v"), c.get("d"), c.get("u"), c.get("f")
    return c, None, None, None


def _write_sheet(ws, sheet, start_row=1, start_col=1, header_note=None):
    cols = sheet["columns"]
    row = start_row

    if header_note:
        c = ws.cell(row, start_col, header_note)
        c.font = TITLE_FONT
        ws.row_dimensions[row].height = 20
        row += 1

    for i, col in enumerate(cols):
        letter = get_column_letter(start_col + i)
        ws.column_dimensions[letter].width = max(6, col["width"] / 7.5)
        cell = ws.cell(row, start_col + i, col["label"])
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal=ALIGN.get(col["align"], "left"),
                                   vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24
    header_row = row
    row += 1

    if not sheet["rows"]:
        ws.cell(row, start_col, sheet.get("empty", "No rows.")).font = SUB_FONT
        return header_row

    for r_i, data_row in enumerate(sheet["rows"]):
        ws.row_dimensions[row].height = 30
        banded = r_i % 2 == 1
        for c_i, raw in enumerate(data_row):
            value, display, url, fmt = _unpack(raw)
            col = cols[c_i]
            cell = ws.cell(row, start_col + c_i)
            cell.value = value if value is not None else (display or "")
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal=ALIGN.get(col["align"], "left"),
                                       vertical="center", wrap_text=col["kind"] == "text")
            cell.border = BOTTOM
            if banded:
                cell.fill = BAND
            if fmt:
                cell.number_format = fmt
            if url:
                cell.hyperlink = url
                cell.font = LINK_FONT
            if col["key"] == "margin":
                cell.font = POS_FONT
        row += 1
    return header_row


def build_workbook(sheets: list) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(title=sheet["name"][:31])
        header_row = _write_sheet(ws, sheet, header_note=sheet.get("hint"))
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        if sheet["rows"]:
            ws.auto_filter.ref = (
                f"A{header_row}:"
                f"{get_column_letter(len(sheet['columns']))}"
                f"{header_row + len(sheet['rows'])}"
            )
    return wb


def workbook_bytes(sheets: list) -> bytes:
    buf = io.BytesIO()
    build_workbook(sheets).save(buf)
    return buf.getvalue()


def save_workbook(sheets: list, path: str) -> str:
    build_workbook(sheets).save(path)
    return path
